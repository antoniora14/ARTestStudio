from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


BASE_GENERATOR = Path(__file__).with_name("generate_phase2_interaction_report.py")
SPEC = spec_from_file_location("phase2_report_helpers", BASE_GENERATOR)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Unable to load report helpers from {BASE_GENERATOR}")
qa = module_from_spec(SPEC)
SPEC.loader.exec_module(qa)


CASES = [
    (
        "MT-P3-001",
        "New diagrams use JSON v2 and round-trip",
        "Verify that new .atd files use the documented JSON v2 schema and reopen without model loss.",
        [
            "Close every ARTestStudio instance.",
            "Launch artifacts/bin/x64/Release/ARTestStudio.exe.",
            "Prepare a writable folder for execution evidence and generated .atd files.",
        ],
        [
            "Create a new diagram with one rectangle and one diamond block.",
            "Connect the rectangle bottom port to the diamond top port.",
            "Use File > Save As and save the file as MT-P3-json-v2.atd.",
            "Keep ARTestStudio open and open MT-P3-json-v2.atd in a text editor.",
            "Verify the root fields format and version, then capture evidence of the JSON.",
            "Close the diagram in ARTestStudio and reopen MT-P3-json-v2.atd.",
            "Compare block types, labels, positions, dimensions, connection endpoints, and route.",
        ],
        [
            "The file is readable UTF-8 JSON and begins with an object, not ARTESTSTUDIO_DIAGRAM 1.",
            "format equals ARTestStudio.Diagram and version equals 2.",
            "nodes and connections are JSON arrays with stable numeric identifiers.",
            "The reopened diagram is visually and functionally equivalent to the saved diagram.",
            "The title remains MT-P3-json-v2.atd and no fault dialog is displayed.",
        ],
    ),
    (
        "MT-P3-002",
        "Legacy v1 compatibility and migration on save",
        "Verify that legacy text diagrams remain loadable and migrate to JSON only after a successful save.",
        [
            "Locate quality/manual-tests/phase-03/fixtures/MT-P3-legacy-v1.atd.",
            "Copy it to the evidence folder as MT-P3-legacy-v1-working.atd; do not modify the fixture.",
            "Confirm in a text editor that the working copy begins with ARTESTSTUDIO_DIAGRAM 1.",
        ],
        [
            "Open MT-P3-legacy-v1-working.atd in ARTestStudio.",
            "Verify the LegacyStart and LegacyWait blocks and their connection are present.",
            "Before saving, inspect the working copy in a text editor and confirm it still begins with ARTESTSTUDIO_DIAGRAM 1.",
            "Without moving any block, press Ctrl+S once.",
            "Open the saved working copy in a text editor and inspect its root fields.",
            "Close and reopen MT-P3-legacy-v1-working.atd in ARTestStudio.",
            "Create one additional block and verify its identifier continues after the migrated maximum by saving and inspecting the JSON.",
        ],
        [
            "The original legacy copy opens with node IDs 42 and 7 and connection ID 77 preserved.",
            "After Ctrl+S, the working copy is JSON with format ARTestStudio.Diagram and version 2.",
            "No data is rewritten merely by opening the untouched fixture.",
            "The migrated JSON reopens with the same two blocks and connection.",
            "New identifiers continue safely after the maximum restored IDs.",
        ],
    ),
    (
        "MT-P3-003",
        "Invalid or future JSON is rejected atomically",
        "Verify explicit errors and preservation of the active diagram when JSON cannot be accepted.",
        [
            "Have a valid diagram open with at least two blocks and one connection.",
            "Locate MT-P3-unsupported-v999.atd and MT-P3-dangling-reference.atd under the phase-03 fixtures folder.",
            "Record the active diagram tab, block count, and connection count.",
        ],
        [
            "Use File > Open and select MT-P3-unsupported-v999.atd.",
            "Record the complete error dialog and dismiss it.",
            "Verify the previously active diagram and its contents remain unchanged.",
            "Use File > Open and select MT-P3-dangling-reference.atd.",
            "Record the complete error dialog and dismiss it.",
            "Again verify the original tab, blocks, connection, selection behavior, and responsiveness.",
        ],
        [
            "Version 999 is rejected with an incompatible/unsupported version message.",
            "The dangling connection is rejected as invalid diagram data.",
            "Neither failed open creates a partially loaded document or changes the active model.",
            "The application remains responsive and the original diagram can still be saved.",
            "Both failures are recorded by the existing storage fault handling.",
        ],
    ),
]


def build(output: Path) -> None:
    doc = Document()
    qa.configure(doc)
    footer = doc.sections[0].footer.paragraphs[0]
    for run in footer.runs:
        run.text = run.text.replace(
            "Phase 2 - InteractionController and Undo/Redo",
            "Phase 3 - Versioned Diagram Persistence",
        )

    kicker = doc.add_paragraph()
    qa.font(kicker.add_run("MANUAL QA REPORT"), 10, qa.BLUE, True)
    doc.add_paragraph("ARTestStudio - Phase 3", style="Title")
    doc.add_paragraph(
        "Versioned .atd JSON v2, legacy v1 compatibility, migration, and atomic validation",
        style="Subtitle",
    )

    values = [
        ("Document ID", "ARTS-QA-P3-JSON-001"),
        ("Version", "1.0"),
        ("Prepared", "August 29, 2026"),
        ("Target build", "Release | x64 | Visual Studio 18 Insiders / v145"),
        ("Execution status", "PENDING MANUAL EXECUTION"),
        ("Tester", "________________________________________"),
    ]
    metadata = doc.add_table(rows=len(values), cols=2)
    for row, (name, value) in zip(metadata.rows, values):
        qa.shade(row.cells[0], qa.LIGHT_GRAY)
        qa.font(row.cells[0].paragraphs[0].add_run(name), color=qa.NAVY, bold=True)
        qa.font(row.cells[1].paragraphs[0].add_run(value))
    qa.table_geometry(metadata, [2160, 7200])

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    qa.shade(callout.cell(0, 0), "E6F4EA")
    paragraph = callout.cell(0, 0).paragraphs[0]
    qa.font(paragraph.add_run("Automated baseline: "), color=qa.NAVY, bold=True)
    qa.font(
        paragraph.add_run(
            "43 of 43 Google Test cases passed in Release x64, including six JSON/migration regressions."
        ),
        color=qa.NAVY,
    )
    qa.table_geometry(callout, [9360])

    qa.heading(doc, "Scope", 1)
    bullets = qa.numbering(doc, "ListBullet")
    for item in (
        "New .atd files are UTF-8 JSON using schema version 2.",
        "Legacy ARTESTSTUDIO_DIAGRAM 1 files remain readable and migrate on the next save.",
        "Future versions, malformed JSON, invalid UTF-8, and dangling references are rejected transactionally.",
        "Existing 16 MB, 64 KB label, model validation, and atomic replacement protections remain active.",
    ):
        qa.list_item(doc, item, bullets)

    qa.heading(doc, "Execution summary", 1)
    summary = doc.add_table(rows=1, cols=4)
    for index, value in enumerate(("Test ID", "Manual test case", "Verdict", "Evidence reference")):
        qa.shade(summary.cell(0, index), qa.NAVY)
        paragraph = summary.cell(0, index).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        qa.font(paragraph.add_run(value), 9.5, qa.WHITE, True)
    header = qa.OxmlElement("w:tblHeader")
    header.set(qa.qn("w:val"), "true")
    summary.rows[0]._tr.get_or_add_trPr().append(header)
    for test_id, title, *_ in CASES:
        cells = summary.add_row().cells
        for index, value in enumerate((test_id, title, "PENDING", "")):
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            qa.font(paragraph.add_run(value), 9.5, qa.GRAY if index == 2 else qa.BLACK, index == 0)
    qa.table_geometry(summary, [1224, 4680, 1296, 2160])

    qa.heading(doc, "Overall verdict", 1)
    qa.label(doc, "Final result", "[ ] PASS     [ ] FAIL     [ ] BLOCKED")
    qa.label(doc, "Defects raised", "____________________________________________________________")
    qa.label(doc, "Approved by", "__________________________    Date: ________________________")

    for test_id, title, objective, preconditions, steps, expected in CASES:
        doc.add_page_break()
        qa.heading(doc, f"{test_id} - {title}", 1)
        qa.label(doc, "Objective", objective, 6)
        qa.label(doc, "Priority", "Critical regression", 8)
        qa.heading(doc, "Preconditions", 2)
        bullets = qa.numbering(doc, "ListBullet")
        for item in preconditions:
            qa.list_item(doc, item, bullets)
        qa.heading(doc, "Exact execution steps", 2)
        numbers = qa.numbering(doc, "ListNumber")
        for item in steps:
            qa.list_item(doc, item, numbers)
        qa.heading(doc, "Expected results / acceptance criteria", 2)
        bullets = qa.numbering(doc, "ListBullet")
        for item in expected:
            qa.list_item(doc, item, bullets)

        doc.add_page_break()
        qa.heading(doc, f"{test_id} - Evidence and execution record", 1)
        qa.label(doc, "Test case", title, 8)
        qa.evidence(doc)

    properties = doc.core_properties
    properties.title = "ARTestStudio Phase 3 JSON Persistence Manual Test Report"
    properties.subject = "Versioned JSON v2, legacy migration, and validation"
    properties.author = "ARTestStudio Quality Assurance"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_phase3_json_persistence_report.py OUTPUT.docx")
    build(Path(sys.argv[1]))
