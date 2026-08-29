from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY, BLUE, DARK_BLUE = "17365D", "2E74B5", "1F4D78"
LIGHT_BLUE, LIGHT_GRAY, GRAY = "E8EEF5", "F2F4F7", "667085"
WHITE, BLACK = "FFFFFF", "000000"

CASES = [
    (
        "MT-P2-001", "Centralized node and connection selection",
        "Verify that a single selection is maintained and is visibly distinguishable.",
        [
            "Launch artifacts/bin/x64/Release/ARTestStudio.exe.",
            "Create a new diagram and drag two blocks from Class View.",
            "Connect both blocks and save the diagram as MT-P2-baseline.atd.",
        ],
        [
            "Left-click the center of the first block without moving the mouse.",
            "Compare the border of the selected block with the second block.",
            "Left-click a blank area of the document.",
            "Right-click directly over an orthogonal segment of the connection line.",
            "Close the context menu without executing a command and inspect the connection.",
            "Right-click the second block.",
        ],
        [
            "Only one element is selected at a time.",
            "A selected block or connection uses a blue, thicker outline; unselected elements remain black.",
            "Clicking blank space clears the selection.",
            "Right-clicking a new element transfers selection without stale highlights.",
        ],
    ),
    (
        "MT-P2-002", "Node drag recorded as one undoable command",
        "Verify state-machine dragging and one-command undo granularity.",
        [
            "Open MT-P2-baseline.atd so the command history starts empty.",
            "Record or capture the initial position of the first block.",
        ],
        [
            "Press and hold the left mouse button at the center of the first block.",
            "Move the pointer through at least three different positions while holding the button.",
            "Release the button at a clearly different final position.",
            "Press Ctrl+Z exactly once.",
            "Open the Edit menu and verify Undo is disabled.",
            "Press Ctrl+Y exactly once.",
        ],
        [
            "The block follows the pointer and connected lines reroute while dragging.",
            "One Ctrl+Z restores the exact initial position; the drag is one command.",
            "With a fresh history, Undo becomes disabled after that single undo.",
            "Ctrl+Y restores the final position and route.",
        ],
    ),
    (
        "MT-P2-003", "Connection creation undo and redo",
        "Verify connection creation through the interaction state machine and command history.",
        [
            "Open MT-P2-baseline.atd and delete its connection if necessary, then save and reopen.",
            "Ensure two unconnected blocks are visible.",
        ],
        [
            "Press the left mouse button over a connection port on the first block.",
            "Move toward a compatible port on the second block and observe the preview.",
            "Release over the target port.",
            "Press Ctrl+Z.",
            "Press Ctrl+Y.",
        ],
        [
            "The preview is a thin blue line that follows the pointer.",
            "Release creates one orthogonal connection and selects it.",
            "Ctrl+Z removes the complete connection.",
            "Ctrl+Y restores the same connection and orthogonal route.",
        ],
    ),
    (
        "MT-P2-004", "Delete node with dependent connections",
        "Verify atomic deletion and restoration of a node plus all dependent connections.",
        [
            "Open a diagram containing at least three blocks.",
            "Connect the middle block to the other two blocks and save/reopen the diagram.",
        ],
        [
            "Right-click the middle block.",
            "Choose Delete from the context menu.",
            "Verify the block and both dependent connections disappear.",
            "Press Ctrl+Z.",
            "Press Ctrl+Y.",
        ],
        [
            "Delete removes the selected block and every connection that references it.",
            "Undo restores the same block, label, geometry, connections, and routes as one operation.",
            "Redo removes the same complete set again without leaving dangling lines.",
        ],
    ),
    (
        "MT-P2-005", "Delete a selected connection only",
        "Verify line hit testing, connection selection, deletion, undo, and redo.",
        [
            "Open a diagram with two blocks connected by a routed line.",
            "Save and reopen the diagram to isolate command history.",
        ],
        [
            "Right-click the middle of a horizontal or vertical connection segment.",
            "Verify the connection is highlighted blue.",
            "Choose Delete from the context menu.",
            "Confirm both blocks remain and only the connection disappears.",
            "Press Ctrl+Z and then Ctrl+Y.",
        ],
        [
            "The line is selectable within a small tolerance around its visible path.",
            "Delete affects only the selected connection.",
            "Undo restores the routed connection; redo removes it again.",
        ],
    ),
    (
        "MT-P2-006", "Copy and paste at the context position",
        "Verify controller-owned clipboard behavior and command history for paste.",
        [
            "Open a diagram containing a diamond block with a recognizable label.",
            "Save and reopen before execution.",
        ],
        [
            "Right-click the diamond and choose Copy.",
            "Right-click a blank area far from existing blocks.",
            "Verify Paste is enabled and choose Paste.",
            "Inspect the new block type, label, size, and center position.",
            "Press Ctrl+Z and then Ctrl+Y.",
        ],
        [
            "Paste creates a new block at the latest right-click position.",
            "The copy preserves type, label, width, and height but receives a new identifier.",
            "The pasted block becomes selected.",
            "Undo removes only the pasted copy; redo restores it.",
        ],
    ),
    (
        "MT-P2-007", "Cut uses one delete command and preserves clipboard",
        "Verify Cut is composed from centralized Copy and Delete operations.",
        [
            "Open a saved diagram with at least one isolated block.",
            "Reopen the file so history is empty.",
        ],
        [
            "Right-click the isolated block and choose Cut.",
            "Press Ctrl+Z once.",
            "Press Ctrl+Y once.",
            "Right-click a blank location and choose Paste.",
        ],
        [
            "Cut removes the selected block and stores its content in the controller clipboard.",
            "One Undo restores the block; one Redo removes it again.",
            "Paste remains available after undo/redo and creates a copy at the context location.",
        ],
    ),
    (
        "MT-P2-008", "New edit invalidates the redo branch",
        "Verify standard linear command-history behavior after undo.",
        ["Open a saved diagram with two blocks and no pending command history."],
        [
            "Move the first block and release it.",
            "Press Ctrl+Z to undo the movement.",
            "Move the second block and release it.",
            "Open the Edit menu and inspect Redo.",
            "Press Ctrl+Y.",
        ],
        [
            "The second movement is accepted as a new command.",
            "Redo is disabled because the former redo branch was abandoned.",
            "Ctrl+Y does not reapply the first block old movement.",
        ],
    ),
    (
        "MT-P2-009", "Orthogonal rerouting remains stable across undo and redo",
        "Verify interaction refactoring did not regress automatic routing.",
        [
            "Create two connected blocks with a third block positioned as an obstacle.",
            "Save and reopen the diagram.",
        ],
        [
            "Move one endpoint block so the connection must choose a visibly different route.",
            "Inspect every segment and bend after releasing the block.",
            "Press Ctrl+Z and inspect the restored route.",
            "Press Ctrl+Y and inspect the reapplied route.",
        ],
        [
            "Every connection segment remains horizontal or vertical with 90-degree bends.",
            "No route crosses the interior of the obstacle block.",
            "Undo restores the previous node position and route together.",
            "Redo reapplies the new position and a valid routed line.",
        ],
    ),
    (
        "MT-P2-010", "Document lifecycle resets transient interaction and history",
        "Verify history is document-scoped and does not leak through New/Open operations.",
        [
            "Have MT-P2-baseline.atd available.",
            "Launch the Release x64 executable.",
        ],
        [
            "Open MT-P2-baseline.atd, move a block, and verify Undo becomes enabled.",
            "Save the diagram to a new .atd file.",
            "Use File > New and open the Edit menu.",
            "Open the saved .atd file and inspect the Edit menu again.",
            "Verify the saved block position and connections are present.",
            "Press Ctrl+Z and Ctrl+Y in the freshly opened document.",
        ],
        [
            "New and Open clear selection, pointer state, undo history, and redo history.",
            "Undo and Redo are disabled immediately after New/Open.",
            "The saved model is not changed by stale commands from the previous document.",
            "The application remains responsive with no unexpected fault dialog or log entry.",
        ],
    ),
    (
        "MT-P2-011", "Autosave preserves the visible document name",
        "Verify that MFC recovery autosave writes a copy without replacing the active document path or tab title.",
        [
            "Close every running ARTestStudio instance.",
            "Launch artifacts/bin/x64/Release/ARTestStudio.exe and open MT-P2-baseline.atd.",
            "Confirm that the tab initially displays exactly MT-P2-baseline.atd.",
        ],
        [
            "Move one block and release it so the document becomes modified.",
            "Record the complete tab title before autosave.",
            "Keep ARTestStudio open for 6 minutes without using Save; the default MFC autosave interval is 5 minutes.",
            "Bring ARTestStudio to the foreground and inspect the complete tab title.",
            "Press Ctrl+S and confirm the tab title again.",
            "Use File > Save As to save as MT-P2-renamed.atd and inspect the tab title.",
        ],
        [
            "Before and after autosave, the tab remains exactly MT-P2-baseline.atd with no hexadecimal prefix.",
            "Autosave does not change the active document path, title, selection, diagram, or responsiveness.",
            "Ctrl+S saves to the original MT-P2-baseline.atd document.",
            "Save As changes the tab exactly to MT-P2-renamed.atd, with no generated prefix or suffix.",
        ],
    ),
]


def font(run, size=11, color=BLACK, bold=False, italic=False):
    run.font.name = "Calibri"
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold, run.italic = bold, italic


def shade(cell, fill):
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(node)


def table_geometry(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = table._tbl.tblPr
    width = pr.find(qn("w:tblW"))
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    indent = OxmlElement("w:tblInd")
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    pr.append(indent)
    grid = OxmlElement("w:tblGrid")
    for value in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(value))
        grid.append(col)
    table._tbl.replace(table._tbl.tblGrid, grid)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            margins = OxmlElement("w:tcMar")
            for name, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
                item = OxmlElement(f"w:{name}")
                item.set(qn("w:w"), str(value))
                item.set(qn("w:type"), "dxa")
                margins.append(item)
            tc_pr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "start", "bottom", "end"):
                border = OxmlElement(f"w:{edge}")
                border.set(qn("w:val"), "single")
                border.set(qn("w:sz"), "6")
                border.set(qn("w:color"), "D0D5DD")
                borders.append(border)
            tc_pr.append(borders)


def numbering(doc, paragraph_style):
    root = doc.part.numbering_part.element
    abstract_id = None
    for abstract in root.findall(qn("w:abstractNum")):
        for level in abstract.findall(qn("w:lvl")):
            style = level.find(qn("w:pStyle"))
            if style is not None and style.get(qn("w:val")) == paragraph_style:
                abstract_id = int(abstract.get(qn("w:abstractNumId")))
                break
        if abstract_id is not None:
            break
    if abstract_id is None:
        raise RuntimeError(f"Numbering style not found: {paragraph_style}")
    num_id = max([int(x.get(qn("w:numId"))) for x in root.findall(qn("w:num"))], default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abstract_id))
    num.append(ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    root.append(num)
    return num_id


def list_item(doc, value, num_id):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    num_pr = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([level, num])
    p._p.get_or_add_pPr().append(num_pr)
    font(p.add_run(value))


def heading(doc, value, level):
    p = doc.add_paragraph(value, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.left_indent = Inches(0)
    p.paragraph_format.first_line_indent = Inches(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def label(doc, name, value, after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    font(p.add_run(f"{name}: "), color=NAVY, bold=True)
    font(p.add_run(value))


def configure(doc):
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.75)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Calibri", Pt(11)
    normal.paragraph_format.space_after, normal.paragraph_format.line_spacing = Pt(6), 1.25
    for name, size, color, before, after in (
        ("Title", 26, NAVY, 0, 8),
        ("Subtitle", 13, GRAY, 0, 16),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name, style.font.size = "Calibri", Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before, style.paragraph_format.space_after = Pt(before), Pt(after)
    header = section.header.paragraphs[0]
    font(header.add_run("ARTestStudio  |  Quality Assurance"), 9, GRAY, True)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    font(footer.add_run("Phase 2 - InteractionController and Undo/Redo    |    Page "), 9, GRAY)
    run = footer.add_run()
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = " PAGE "
    separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t"); text.text = "1"
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def metadata(doc):
    values = [
        ("Document ID", "ARTS-QA-P2-IC-001"),
        ("Version", "1.0"),
        ("Prepared", "August 28, 2026"),
        ("Target build", "Release | x64 | Visual Studio 18 Insiders / v145"),
        ("Execution status", "PENDING MANUAL EXECUTION"),
        ("Tester", "________________________________________"),
    ]
    table = doc.add_table(rows=len(values), cols=2)
    for row, (name, value) in zip(table.rows, values):
        shade(row.cells[0], LIGHT_GRAY)
        font(row.cells[0].paragraphs[0].add_run(name), color=NAVY, bold=True)
        font(row.cells[1].paragraphs[0].add_run(value))
    table_geometry(table, [2160, 7200])


def evidence(doc):
    table = doc.add_table(rows=2, cols=1)
    shade(table.cell(0, 0), LIGHT_GRAY)
    font(table.cell(0, 0).paragraphs[0].add_run(
        "Paste screenshots here. Include the complete document window and, when relevant, "
        "the Edit menu showing Undo/Redo enabled or disabled."
    ), 9.5, GRAY, italic=True)
    table.rows[1].height = Inches(5.1)
    table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    font(table.cell(1, 0).paragraphs[0].add_run(
        "Evidence reference / filename: ______________________________________________"
    ), 9.5, GRAY)
    table_geometry(table, [9360])
    doc.add_paragraph()
    label(doc, "Actual result", "_______________________________________________________________", 2)
    label(doc, "Verdict", "[ ] PASS     [ ] FAIL     [ ] BLOCKED", 2)
    label(doc, "Executed by", "________________________", 2)
    label(doc, "Date / build", "________________________", 2)
    label(doc, "Notes / defect ID", "________________________________________________________", 0)


def build(output):
    doc = Document()
    configure(doc)
    kicker = doc.add_paragraph()
    font(kicker.add_run("MANUAL QA REPORT"), 10, BLUE, True)
    doc.add_paragraph("ARTestStudio - Phase 2", style="Title")
    doc.add_paragraph(
        "InteractionController, centralized selection, Command/Memento history, Undo and Redo",
        style="Subtitle",
    )
    metadata(doc)
    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    shade(callout.cell(0, 0), "E6F4EA")
    p = callout.cell(0, 0).paragraphs[0]
    font(p.add_run("Automated baseline: "), color=NAVY, bold=True)
    font(p.add_run(
        "37 of 37 Google Test cases passed in Release x64. Reports are under "
        "artifacts/test-results/x64/Release."
    ), color=NAVY)
    table_geometry(callout, [9360])
    heading(doc, "Scope", 1)
    scope_bullets = numbering(doc, "ListBullet")
    for item in (
        "Centralized node/connection selection and selected-state rendering.",
        "State machine for idle, node dragging, and connection creation.",
        "Command + Memento history for add, move, connect, delete, cut, paste, undo, and redo.",
        "Automatic orthogonal rerouting and document lifecycle regression.",
    ):
        list_item(doc, item, scope_bullets)
    heading(doc, "Execution rule", 1)
    doc.add_paragraph(
        "Execute the cases in order. Start from the stated preconditions, capture evidence before moving "
        "to the next case, and mark PASS only when every expected result is observed."
    )

    doc.add_page_break()
    heading(doc, "Execution summary", 1)
    doc.add_paragraph(
        "Complete this matrix after every case. Reference screenshots, videos, or defect records."
    )
    summary = doc.add_table(rows=1, cols=4)
    for index, value in enumerate(("Test ID", "Manual test case", "Verdict", "Evidence reference")):
        shade(summary.cell(0, index), NAVY)
        p = summary.cell(0, index).paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run(value), 9.5, WHITE, True)
    header = OxmlElement("w:tblHeader"); header.set(qn("w:val"), "true")
    summary.rows[0]._tr.get_or_add_trPr().append(header)
    for test_id, title, *_ in CASES:
        cells = summary.add_row().cells
        for index, value in enumerate((test_id, title, "PENDING", "")):
            p = cells[index].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            font(p.add_run(value), 9.5, GRAY if index == 2 else BLACK, index == 0)
    table_geometry(summary, [1224, 4680, 1296, 2160])
    heading(doc, "Overall verdict", 1)
    label(doc, "Final result", "[ ] PASS     [ ] FAIL     [ ] BLOCKED")
    label(doc, "Defects raised", "____________________________________________________________")
    label(doc, "Approved by", "__________________________    Date: ________________________")
    label(doc, "Comments", "________________________________________________________________")

    for test_id, title, objective, preconditions, steps, expected in CASES:
        doc.add_page_break()
        heading(doc, f"{test_id} - {title}", 1)
        label(doc, "Objective", objective, 6)
        label(doc, "Priority", "Critical regression", 8)
        heading(doc, "Preconditions", 2)
        precondition_bullets = numbering(doc, "ListBullet")
        for item in preconditions:
            list_item(doc, item, precondition_bullets)
        heading(doc, "Exact execution steps", 2)
        step_numbers = numbering(doc, "ListNumber")
        for item in steps:
            list_item(doc, item, step_numbers)
        heading(doc, "Expected results / acceptance criteria", 2)
        result_bullets = numbering(doc, "ListBullet")
        for item in expected:
            list_item(doc, item, result_bullets)
        doc.add_page_break()
        heading(doc, f"{test_id} - Evidence and execution record", 1)
        label(doc, "Test case", title, 8)
        evidence(doc)

    props = doc.core_properties
    props.title = "ARTestStudio Phase 2 Manual Test Report"
    props.subject = "InteractionController, selection, commands, undo and redo"
    props.author = "ARTestStudio Quality Assurance"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_interaction_report.py OUTPUT.docx")
    build(Path(sys.argv[1]))
