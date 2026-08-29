from importlib.util import module_from_spec, spec_from_file_location
from pathlib import Path
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH


BASE_GENERATOR = Path(__file__).with_name("generate_phase2_interaction_report.py")
SPEC = spec_from_file_location("phase2_report", BASE_GENERATOR)
if SPEC is None or SPEC.loader is None:
    raise RuntimeError(f"Unable to load report helpers from {BASE_GENERATOR}")
qa = module_from_spec(SPEC)
SPEC.loader.exec_module(qa)

CASES = [case for case in qa.CASES if case[0] == "MT-P2-011"]


def build(output: Path) -> None:
    doc = Document()
    qa.configure(doc)

    kicker = doc.add_paragraph()
    qa.font(kicker.add_run("MANUAL QA REGRESSION ADDENDUM"), 10, qa.BLUE, True)
    doc.add_paragraph("ARTestStudio - Phase 2", style="Title")
    doc.add_paragraph(
        "Document title stability during MFC recovery autosave",
        style="Subtitle",
    )

    values = [
        ("Document ID", "ARTS-QA-P2-IC-ADD-001"),
        ("Version", "1.0"),
        ("Prepared", "August 29, 2026"),
        ("Target build", "Release | x64 | Visual Studio 18 Insiders / v145"),
        ("Execution status", "PENDING MANUAL EXECUTION"),
        ("Tester", "________________________________________"),
    ]
    metadata = doc.add_table(rows=len(values), cols=2)
    for row, (name, value) in zip(metadata.rows, values):
        qa.shade(row.cells[0], qa.LIGHT_GRAY)
        qa.font(row.cells[0].paragraphs[0].add_run(name), color=qa.NAVY, bold=True)
        qa.font(row.cells[1].paragraphs[0].add_run(value))
    qa.table_geometry(metadata, [2160, 7200])

    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    qa.shade(callout.cell(0, 0), "E6F4EA")
    paragraph = callout.cell(0, 0).paragraphs[0]
    qa.font(paragraph.add_run("Automated baseline: "), color=qa.NAVY, bold=True)
    qa.font(
        paragraph.add_run(
            "37 of 37 Google Test cases passed in Debug and Release x64 after the correction."
        ),
        color=qa.NAVY,
    )
    qa.table_geometry(callout, [9360])

    qa.heading(doc, "Defect and correction scope", 1)
    doc.add_paragraph(
        "Regression for the defect where MFC recovery autosave replaced the visible tab title "
        "with a generated hexadecimal autosave filename. The correction keeps path ownership "
        "inside CDocument::DoSave and preserves normal Save As behavior."
    )

    qa.heading(doc, "Execution summary", 1)
    summary = doc.add_table(rows=1, cols=4)
    for index, value in enumerate(("Test ID", "Manual test case", "Verdict", "Evidence reference")):
        qa.shade(summary.cell(0, index), qa.NAVY)
        paragraph = summary.cell(0, index).paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        qa.font(paragraph.add_run(value), 9.5, qa.WHITE, True)
    for test_id, title, *_ in CASES:
        cells = summary.add_row().cells
        for index, value in enumerate((test_id, title, "PENDING", "")):
            paragraph = cells[index].paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 1 else WD_ALIGN_PARAGRAPH.CENTER
            qa.font(paragraph.add_run(value), 9.5, qa.GRAY if index == 2 else qa.BLACK, index == 0)
    qa.table_geometry(summary, [1224, 4680, 1296, 2160])

    for test_id, title, objective, preconditions, steps, expected in CASES:
        doc.add_page_break()
        qa.heading(doc, f"{test_id} - {title}", 1)
        qa.label(doc, "Objective", objective, 6)
        qa.label(doc, "Priority", "Critical regression", 8)
        qa.heading(doc, "Preconditions", 2)
        bullets = qa.numbering(doc, "ListBullet")
        for item in preconditions:
            qa.list_item(doc, item, bullets)
        qa.heading(doc, "Exact execution steps", 2)
        numbers = qa.numbering(doc, "ListNumber")
        for item in steps:
            qa.list_item(doc, item, numbers)
        qa.heading(doc, "Expected results / acceptance criteria", 2)
        bullets = qa.numbering(doc, "ListBullet")
        for item in expected:
            qa.list_item(doc, item, bullets)

        doc.add_page_break()
        qa.heading(doc, f"{test_id} - Evidence and execution record", 1)
        qa.label(doc, "Test case", title, 8)
        qa.evidence(doc)

    properties = doc.core_properties
    properties.title = "ARTestStudio Phase 2 Autosave Title Regression Report"
    properties.subject = "MFC recovery autosave must preserve the active document title"
    properties.author = "ARTestStudio Quality Assurance"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage: generate_autosave_title_report.py OUTPUT.docx")
    build(Path(sys.argv[1]))
