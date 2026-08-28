from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
GREEN = "E8F3EC"
GOLD = "FFF4D6"


def set_run_font(run, name="Calibri", size=11, color=INK, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, name, size, color):
    style.font.name = name
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, "Calibri", 11, INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for style_name, (size, color, before, after) in tokens.items():
        style = doc.styles[style_name]
        set_style_font(style, "Calibri", size, color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
    for style_name in ("Header", "Footer"):
        style = doc.styles[style_name]
        set_style_font(style, "Calibri", 8.5, MUTED)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.0


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table, color="C8D0DA", size="4"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=TABLE_INDENT_DXA, borders=True):
    if sum(widths) != CONTENT_WIDTH_DXA and indent == TABLE_INDENT_DXA:
        raise ValueError(f"Body table widths must total {CONTENT_WIDTH_DXA}: {widths}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if borders:
        set_table_borders(table)


def prevent_row_split(row):
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    row._tr.get_or_add_trPr().append(node)


def format_cell(cell, text, *, bold=False, color=INK, size=9, fill=None, align=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    if align is not None:
        paragraph.alignment = align
    run = paragraph.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold)
    if fill:
        set_cell_shading(cell, fill)


def add_fixed_table(doc, widths, rows, header=False, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=0, cols=len(widths))
    set_table_geometry(table, widths)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for col_index, value in enumerate(values):
            is_header = header and row_index == 0
            format_cell(
                row.cells[col_index],
                str(value),
                bold=is_header,
                color=DARK_BLUE if is_header else INK,
                fill=header_fill if is_header else None,
            )
        if header and row_index == 0:
            repeat_header(row)
    return table


def add_spacer(doc, points=4):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(points)
    paragraph.paragraph_format.line_spacing = 1.0


def add_rule(paragraph, color=BLUE, size="12"):
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    paragraph._p.get_or_add_pPr().append(p_bdr)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=8.5, color=MUTED)


def configure_page(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def configure_header_footer(section, document_id):
    header = section.header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = ""
    table = header.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [6480, 2880], indent=0, borders=False)
    format_cell(table.cell(0, 0), "ARTestStudio | Quality Assurance", bold=True, color=DARK_BLUE, size=8.5)
    format_cell(table.cell(0, 1), document_id, bold=True, color=MUTED, size=8.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
    footer = section.footer
    footer.is_linked_to_previous = False
    footer.paragraphs[0].text = ""
    table = footer.add_table(rows=1, cols=2, width=Inches(6.5))
    set_table_geometry(table, [6480, 2880], indent=0, borders=False)
    format_cell(table.cell(0, 0), "Registro controlado de pruebas manuales", color=MUTED, size=8.5)
    right = table.cell(0, 1)
    right.text = ""
    paragraph = right.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("Pagina ")
    set_run_font(run, size=8.5, color=MUTED)
    add_page_field(paragraph)


def add_label_value_table(doc, pairs, fill=LIGHT_GRAY):
    table = add_fixed_table(doc, [1875, 7485], pairs)
    for row in table.rows:
        format_cell(row.cells[0], row.cells[0].text, bold=True, color=DARK_BLUE, fill=fill)
    return table


def add_callout(doc, title, body, fill=GOLD):
    table = add_fixed_table(doc, [9360], [("",)])
    set_cell_shading(table.cell(0, 0), fill)
    paragraph = table.cell(0, 0).paragraphs[0]
    title_run = paragraph.add_run(title + "\n")
    set_run_font(title_run, size=9.5, color=DARK_BLUE, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, size=9.5, color=INK)


def add_test_case(doc, case):
    doc.add_page_break()
    doc.add_heading(f"{case['id']} - {case['title']}", level=1)
    add_label_value_table(doc, [
        ("Prioridad", case["priority"]),
        ("Tipo", case["type"]),
        ("Requisitos", case["requirements"]),
        ("Cobertura automatica", case["automation"]),
    ])
    doc.add_heading("Objetivo", level=2)
    paragraph = doc.add_paragraph(case["objective"])
    paragraph.paragraph_format.keep_with_next = True
    doc.add_heading("Precondiciones y datos", level=2)
    table = add_fixed_table(doc, [1875, 7485], case["preconditions"])
    for row in table.rows:
        format_cell(row.cells[0], row.cells[0].text, bold=True, color=DARK_BLUE, fill=LIGHT_GRAY)
    doc.add_heading("Procedimiento de ejecucion", level=2)
    rows = [("Paso", "Accion exacta", "Resultado esperado")]
    rows.extend((str(index), action, expected) for index, (action, expected) in enumerate(case["steps"], start=1))
    table = add_fixed_table(doc, [620, 4400, 4340], rows, header=True)
    for row in table.rows[1:]:
        format_cell(row.cells[0], row.cells[0].text, bold=True, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_heading("Registro de ejecucion", level=2)
    table = add_fixed_table(doc, [1875, 7485], [
        ("Estado", "[ ] PASS    [ ] FAIL    [ ] BLOCKED    [ ] NOT RUN"),
        ("Ejecutor", "____________________________________________"),
        ("Fecha / hora", "____________________________________________"),
        ("Resultado real", "\n\n"),
        ("Defecto asociado", "ID: ____________________   Severidad: ____________________"),
    ])
    for row in table.rows:
        format_cell(row.cells[0], row.cells[0].text, bold=True, color=DARK_BLUE, fill=LIGHT_GRAY)
    if case["id"] == "TC-PH-006":
        doc.add_page_break()
    doc.add_heading("Evidencia", level=2)
    add_fixed_table(doc, [1900, 2800, 4660], [
        ("Evidencia ID", "Archivo / captura", "Descripcion y punto verificado"),
        (f"EV-{case['id']}-01", "", ""),
        (f"EV-{case['id']}-02", "", ""),
        (f"EV-{case['id']}-03", "", ""),
    ], header=True)
    if case["id"] == "TC-PH-006":
        note = doc.add_paragraph(
            "Pagina reservada para las capturas del error de reemplazo, hashes SHA256 antes/despues y verificacion del diagrama conservado."
        )
        note.paragraph_format.space_before = Pt(6)
        for run in note.runs:
            set_run_font(run, size=9, color=MUTED, italic=True)


def audit_document(doc):
    for section in doc.sections:
        assert section.page_width == Inches(8.5)
        assert section.page_height == Inches(11)
        assert section.left_margin == Inches(1)
        assert section.right_margin == Inches(1)
    for table in doc.tables:
        layout = table._tbl.tblPr.find(qn("w:tblLayout"))
        assert layout is not None and layout.get(qn("w:type")) == "fixed"
        tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
        assert tbl_w is not None and int(tbl_w.get(qn("w:w"))) == CONTENT_WIDTH_DXA
        assert sum(int(col.get(qn("w:w"))) for col in table._tbl.tblGrid) == CONTENT_WIDTH_DXA


def build_report(output_path: Path):
    document_id = "ARTEST-QA-MTR-PH-001"
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_page(section)
        configure_header_footer(section, document_id)

    add_spacer(doc, 16)
    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(4)
    set_run_font(kicker.add_run("MANUAL TEST REPORT"), size=10, color=BLUE, bold=True)
    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    set_run_font(title.add_run("ARTestStudio"), size=27, color=INK, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_run_font(subtitle.add_run("Phase 2 - Persistence Hardening and Recovery"), size=15, color=DARK_BLUE, bold=True)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(12)
    add_rule(rule)

    add_label_value_table(doc, [
        ("Document ID", document_id),
        ("Version", "1.0"),
        ("Status", "READY FOR MANUAL EXECUTION"),
        ("Prepared", "2026-08-27"),
        ("Base commit", "2cbce34 (feature under test is currently uncommitted)"),
        ("Repository", r"D:\GitHub\main\ARTestStudio"),
        ("Release executable", r"artifacts\bin\x64\Release\ARTestStudio.exe"),
    ])
    add_spacer(doc, 8)
    add_callout(doc, "Release gate",
        "All seven manual cases must be PASS. Every executed case requires at least one referenced evidence item. "
        "Any failed High-priority case blocks closure of this persistence increment.", fill=GREEN)

    doc.add_heading("Document control", level=1)
    add_fixed_table(doc, [1200, 1400, 1760, 5000], [
        ("Version", "Date", "Author", "Change"),
        ("1.0", "2026-08-27", "Codex / ARTestStudio", "Initial controlled report for persistence hardening."),
    ], header=True)
    doc.add_heading("Scope and acceptance criteria", level=1)
    table = add_fixed_table(doc, [2100, 7260], [
        ("In scope", "16 MB file limit; 64 KB UTF-8 label limit; bounded parsing; transactional load; atomic save; stale .tmp cleanup; replacement-failure recovery; structured logs."),
        ("Out of scope", "Project files .atprj, cloud synchronization, multi-user file locking and line editing interactions."),
        ("Pass rule", "Observed behavior equals every expected result; the active or previous diagram remains intact after each negative test."),
        ("Evidence rule", "Use the Evidence ID from each case in screenshot names, attachments and defect references."),
    ])
    for row in table.rows:
        format_cell(row.cells[0], row.cells[0].text, bold=True, color=DARK_BLUE, fill=LIGHT_GRAY)
    doc.add_heading("Automated regression already executed", level=1)
    add_fixed_table(doc, [1300, 1300, 1400, 1560, 3800], [
        ("Configuration", "Platform", "Result", "Date", "Command"),
        ("Debug", "x64", "29 / 29 PASS", "2026-08-27", r"scripts\build.cmd -Configuration Debug"),
        ("Release", "x64", "29 / 29 PASS", "2026-08-27", r"scripts\build.cmd -Configuration Release"),
    ], header=True)
    doc.add_heading("Manual execution environment", level=1)
    add_label_value_table(doc, [
        ("Tester", "____________________________________________"),
        ("Execution date", "____________________________________________"),
        ("Windows version", "____________________________________________"),
        ("Visual Studio", r"18 Insiders - D:\Program Files\Microsoft Visual Studio\18\Insiders"),
        ("Configuration", "Release | x64"),
        ("Repository root", r"D:\GitHub\main\ARTestStudio"),
        ("Log path", r"%LOCALAPPDATA%\ARTestStudio\Logs\ARTestStudio.log"),
    ])
    doc.add_heading("Preparation procedure", level=1)
    add_fixed_table(doc, [620, 4600, 4140], [
        ("Step", "Action", "Expected result"),
        ("1", r"Open Command Prompt and run: cd /d D:\GitHub\main\ARTestStudio", "The prompt changes to the repository root."),
        ("2", r"Run: scripts\build.cmd -Configuration Release", "Build completes and prints 29/29 tests passed."),
        ("3", r"Run: scripts\manual-tests\generate_persistence_hardening_fixtures.cmd", "A paused window lists four generated fixture paths."),
        ("4", r"Start: artifacts\bin\x64\Release\ARTestStudio.exe", "ARTestStudio opens without an error dialog."),
        ("5", r"Keep File Explorer open at artifacts\manual-tests\persistence-hardening.", "All fixtures are visible; file extensions are displayed."),
    ], header=True)

    cases = [
        {
            "id": "TC-PH-001",
            "title": "Nominal save and load regression",
            "priority": "High",
            "type": "Functional / regression",
            "requirements": "PH-REQ-003, PH-REQ-005",
            "automation": "Yes - round-trip and stable identifiers",
            "objective": "Confirm that normal .atd persistence remains functional after introducing limits and atomic writing.",
            "preconditions": [
                ("Application", "Release executable is running."),
                ("Test file", r"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-baseline.atd"),
                ("Initial state", "No document named PH-baseline.atd exists."),
            ],
            "steps": [
                ("Press Ctrl+N. Drag one Rectangle and one Diamond from Class View to the document.", "A clean diagram contains exactly two visible blocks."),
                ("Connect the Rectangle to the Diamond and visually confirm the orthogonal route.", "One connection is visible and remains attached to both blocks."),
                (r"Use File > Save As and enter D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-baseline.atd.", "Save completes without a dialog; the title reflects PH-baseline.atd."),
                ("Close only the document, then press Ctrl+O and select PH-baseline.atd.", "The file opens without warnings."),
                ("Compare the reopened diagram with the pre-save view.", "Both blocks, their kinds, positions and the connection are preserved."),
            ],
        },
        {
            "id": "TC-PH-002",
            "title": "Reject a diagram larger than 16 MB",
            "priority": "High",
            "type": "Security / negative",
            "requirements": "PH-REQ-001, PH-REQ-005, PH-REQ-006",
            "automation": "Yes - 16 MB + 1 byte boundary",
            "objective": "Verify that total file size is rejected before parsing and that the current valid document is not replaced.",
            "preconditions": [
                ("Active document", "PH-baseline.atd is open and visibly contains two blocks."),
                ("Fixture", r"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-oversized-16mb.atd"),
                ("Exact size", "16,777,217 bytes (16 MB + 1 byte)."),
            ],
            "steps": [
                ("Capture the current PH-baseline.atd window as before-state evidence.", "The capture clearly shows both blocks and their connection."),
                ("Press Ctrl+O and select PH-oversized-16mb.atd.", "An error dialog appears; ARTestStudio does not hang or close."),
                ("Verify the dialog includes: El archivo excede el limite permitido de 16 MB.", "The file-size-specific error is visible."),
                ("Dismiss the dialog and inspect the open documents.", "PH-baseline.atd remains open and unchanged; the oversized fixture is not opened."),
                (r"Open %LOCALAPPDATA%\ARTestStudio\Logs\ARTestStudio.log and search for STORAGE_10.", "A Storage warning records operation abrir and the oversized fixture path."),
            ],
        },
        {
            "id": "TC-PH-003",
            "title": "Reject a label larger than 64 KB",
            "priority": "High",
            "type": "Security / boundary",
            "requirements": "PH-REQ-002, PH-REQ-005, PH-REQ-006",
            "automation": "Yes - bounded quoted-string parser",
            "objective": "Verify bounded label parsing and transactional rejection of a valid-size file containing a label over the per-label limit.",
            "preconditions": [
                ("Active document", "PH-baseline.atd remains open and unchanged."),
                ("Fixture", r"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-label-over-64kb.atd"),
                ("Label size", "65,537 UTF-8 bytes (64 KB + 1 byte)."),
            ],
            "steps": [
                ("Press Ctrl+O and select PH-label-over-64kb.atd.", "An error dialog appears promptly; the application remains responsive."),
                ("Verify the dialog includes: El diagrama excede uno de los limites de seguridad permitidos.", "The limit-specific category is shown."),
                ("Verify the detail includes: Una etiqueta excede el limite de 64 KB.", "The exact violated limit is identified."),
                ("Dismiss the dialog and inspect PH-baseline.atd.", "The baseline diagram is still open and unchanged."),
                (r"Search ARTestStudio.log for STORAGE_11 and PH-label-over-64kb.atd.", "The rejected path and operation abrir are recorded."),
            ],
        },
        {
            "id": "TC-PH-004",
            "title": "Reject a truncated quoted label",
            "priority": "High",
            "type": "Robustness / negative",
            "requirements": "PH-REQ-002, PH-REQ-005, PH-REQ-006",
            "automation": "Yes - truncated label regression",
            "objective": "Confirm that a truncated quoted field is detected without partially loading or damaging an existing document.",
            "preconditions": [
                ("Active document", "PH-baseline.atd remains open and unchanged."),
                ("Fixture", r"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-truncated.atd"),
                ("Fixture condition", "The NODE label starts with a quote but has no closing quote or END section."),
            ],
            "steps": [
                ("Press Ctrl+O and select PH-truncated.atd.", "An invalid-data error dialog appears; no crash occurs."),
                ("Verify the detail includes: La etiqueta de un bloque esta truncada o es invalida.", "The truncated-label cause is explicit."),
                ("Dismiss the dialog.", "No document is created for PH-truncated.atd."),
                ("Inspect PH-baseline.atd.", "The baseline diagram remains intact and editable."),
                (r"Search ARTestStudio.log for STORAGE_8 and PH-truncated.atd.", "A structured Storage warning records the failed open."),
            ],
        },
        {
            "id": "TC-PH-005",
            "title": "Clean a stale .tmp during a successful save",
            "priority": "High",
            "type": "Recovery / filesystem",
            "requirements": "PH-REQ-003, PH-REQ-004",
            "automation": "Yes - real Windows atomic writer",
            "objective": "Verify deterministic cleanup of an abandoned temporary file before writing a new valid destination.",
            "preconditions": [
                ("Temporary fixture", r"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-stale-temporary.atd.tmp"),
                ("Expected content", "The .tmp exists and contains STALE_TEMPORARY_CONTENT."),
                ("Destination", "PH-stale-temporary.atd does not exist before execution."),
            ],
            "steps": [
                ("In File Explorer, confirm PH-stale-temporary.atd.tmp exists; capture its name and size.", "The temporary file is 23 bytes and the final .atd is absent."),
                ("In ARTestStudio press Ctrl+N and drag one Rectangle into the document.", "A new diagram with one block is visible."),
                (r"Use File > Save As and save exactly as D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-stale-temporary.atd.", "Save completes without an error dialog."),
                ("Refresh File Explorer.", "PH-stale-temporary.atd exists and PH-stale-temporary.atd.tmp no longer exists."),
                ("Close the document, press Ctrl+O and open PH-stale-temporary.atd.", "The file opens and contains the saved Rectangle."),
            ],
        },
        {
            "id": "TC-PH-006",
            "title": "Preserve the previous document when replacement fails",
            "priority": "Critical",
            "type": "Recovery / destructive-failure prevention",
            "requirements": "PH-REQ-003, PH-REQ-006",
            "automation": "Yes - injected replacement failure and byte comparison",
            "objective": "Prove that a failed destination replacement never corrupts or overwrites the last valid .atd.",
            "preconditions": [
                ("Baseline file", r"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-lock-preservation.atd"),
                ("Baseline content", "Exactly two blocks and one connection; file is saved."),
                ("Lock helper", r"scripts\manual-tests\hold_file_lock.cmd"),
            ],
            "steps": [
                ("Create the baseline diagram with exactly two blocks and one connection. Save it as PH-lock-preservation.atd.", "The baseline saves and reopens successfully."),
                (r"In PowerShell record: Get-FileHash 'D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-lock-preservation.atd' -Algorithm SHA256", "A SHA256 value is displayed; copy it to evidence."),
                (r"Run: scripts\manual-tests\hold_file_lock.cmd \"D:\GitHub\main\ARTestStudio\artifacts\manual-tests\persistence-hardening\PH-lock-preservation.atd\"", "The helper reports Archivo bloqueado and waits for ENTER."),
                ("Without closing the helper, add a third block in ARTestStudio and press Ctrl+S.", "A replacement error appears; the dialog states that the previous content was preserved."),
                ("Capture the error. Return to the helper window and press ENTER to release the lock.", "The helper prints Bloqueo liberado."),
                ("Close the modified document and choose No if prompted to save. Recompute SHA256.", "The post-failure hash exactly matches the baseline hash."),
                ("Open PH-lock-preservation.atd.", "Only the original two blocks and one connection are present; no .tmp remains."),
                (r"Search ARTestStudio.log for STORAGE_13 and PH-lock-preservation.atd.", "A ReplacementFailure warning records operation guardar."),
            ],
        },
        {
            "id": "TC-PH-007",
            "title": "Recover normal saving after a replacement failure",
            "priority": "High",
            "type": "Recovery / continuity",
            "requirements": "PH-REQ-003, PH-REQ-004, PH-REQ-006",
            "automation": "Partial - successful save after failure paths",
            "objective": "Confirm that releasing the external lock is sufficient for the next save and that the application requires no restart or manual cleanup.",
            "preconditions": [
                ("Previous case", "TC-PH-006 completed and the file lock is released."),
                ("Open document", "PH-lock-preservation.atd is open with its original two-block baseline."),
                ("Temporary state", "PH-lock-preservation.atd.tmp does not exist."),
            ],
            "steps": [
                ("Add one third block to the reopened baseline.", "The document is marked modified and shows three blocks."),
                ("Press Ctrl+S.", "Save completes without an error dialog and the modified flag clears."),
                ("Close and reopen PH-lock-preservation.atd.", "All three blocks are present."),
                ("Refresh File Explorer and inspect the directory.", "No PH-lock-preservation.atd.tmp file remains."),
                ("Review ARTestStudio.log around the prior failure.", "The earlier STORAGE_13 entry remains available; no new Storage warning corresponds to the successful save."),
            ],
        },
    ]

    doc.add_heading("Manual execution summary", level=1)
    rows = [("Test case", "Priority", "Purpose", "Status", "Evidence", "Defect")]
    rows.extend((case["id"], case["priority"], case["title"], "NOT RUN", "", "") for case in cases)
    add_fixed_table(doc, [1250, 1050, 3440, 1100, 1450, 1070], rows, header=True)
    for case in cases:
        add_test_case(doc, case)

    doc.add_page_break()
    doc.add_heading("Final approval", level=1)
    add_callout(doc, "Closure decision",
        "Mark APPROVED only when TC-PH-001 through TC-PH-007 are PASS and all evidence references resolve.",
        fill=GREEN)
    add_spacer(doc, 8)
    table = add_fixed_table(doc, [1875, 7485], [
        ("Overall result", "[ ] APPROVED    [ ] REJECTED    [ ] CONDITIONAL"),
        ("Open defects", "____________________________________________________________"),
        ("QA executor", "Name: ____________________   Signature: ____________________"),
        ("Technical reviewer", "Name: ____________________   Signature: ____________________"),
        ("Approval date", "____________________"),
        ("Comments", "\n\n\n"),
    ])
    for row in table.rows:
        format_cell(row.cells[0], row.cells[0].text, bold=True, color=DARK_BLUE, fill=LIGHT_GRAY)

    audit_document(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    print(f"Created {output_path}")
    print(f"Body tables audited: {len(doc.tables)}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build_report(args.output.resolve())


if __name__ == "__main__":
    main()
